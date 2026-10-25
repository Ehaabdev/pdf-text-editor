"""Download as…: the edited PDF as pictures, PowerPoint, Word or Excel."""

import io
import json
import zipfile
from pathlib import Path
from urllib.parse import unquote

import pymupdf
import pytest
from PIL import Image
from pptx import Presentation

from app import exporting, limits
from conftest import make_pdf, upload

FIXTURES = Path(__file__).parent / "fixtures"
REPORT = (FIXTURES / "images" / "images-word.pdf").read_bytes()


def test_one_page_is_one_picture():
    out = exporting.export(make_pdf(1), "notes.pdf", "png")
    assert out.filename == "notes.png" and out.media_type == "image/png"
    png = Image.open(io.BytesIO(out.data))
    assert png.width == round(595 * 200 / 72)  # 200 dots per inch
    assert round(png.info["dpi"][0]) == 200  # so it prints at the page's size


def test_several_pages_are_pictures_in_a_zip():
    out = exporting.export(make_pdf(3), "notes.pdf", "jpg")
    assert out.filename == "notes (pictures).zip"
    names = zipfile.ZipFile(io.BytesIO(out.data)).namelist()
    assert names == ["notes - page 1.jpg", "notes - page 2.jpg", "notes - page 3.jpg"]


def test_a_huge_page_is_drawn_smaller():
    doc = pymupdf.open()
    doc.new_page(width=5000, height=5000)
    png = Image.open(io.BytesIO(exporting.export(doc.tobytes(), "poster.pdf", "png").data))
    assert png.width * png.height <= exporting.MAX_PAGE_PIXELS * 1.01


def test_powerpoint_as_pictures_has_a_slide_per_page():
    out = exporting.export(REPORT, "report.pdf", "pptx-pictures")
    assert out.filename == "report.pptx"
    deck = Presentation(io.BytesIO(out.data))
    assert len(deck.slides) == 2
    assert abs(deck.slide_width - 595 * 12700) < 12700 and abs(deck.slide_height - 842 * 12700) < 12700
    for slide in deck.slides:
        (picture,) = slide.shapes
        assert picture.shape_type == 13  # a picture
        assert picture.left == 0 and picture.top == 0 and picture.width == deck.slide_width


def test_pages_of_another_size_fit_the_slides():
    doc = pymupdf.open()
    doc.new_page(width=595, height=842)  # A4, upright
    doc.new_page(width=842, height=595)  # A4, sideways
    deck = Presentation(io.BytesIO(exporting.export(doc.tobytes(), "mixed.pdf", "pptx-pictures").data))
    sideways = deck.slides[1].shapes[0]
    assert sideways.width == deck.slide_width  # as wide as the slide, centered top to bottom
    assert sideways.height < deck.slide_height and abs(sideways.top * 2 + sideways.height - deck.slide_height) <= 2


def test_slides_stay_within_powerpoints_limits():
    doc = pymupdf.open()
    doc.new_page(width=6000, height=3000)
    w, h = exporting.slide_size(doc)
    assert w <= 56 * 72 and abs(w / h - 2) < 0.01


def test_download_as_through_the_website(client):
    doc = upload(client, REPORT, "تقرير.pdf").json()
    r = client.get(f"/api/documents/{doc['id']}/export?format=pptx-pictures")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(exporting.PPTX)
    disposition = r.headers["content-disposition"]
    assert disposition.startswith("attachment; filename*=utf-8''") and unquote(disposition.split("''")[1]) == "تقرير.pptx"
    assert client.get(f"/api/documents/{doc['id']}/export?format=exe").status_code == 400


def test_downloads_have_their_own_limit():
    assert limits.kind_of("GET", "/api/documents/abc/export") == "export"
    assert "export" in limits.HEAVY
    assert limits.kind_of("GET", "/api/documents/abc/pages/0/image") == "read"


@pytest.mark.parametrize("name, expected", [("a.pdf", "a"), ("A.PDF", "A"), ("report", "report"), ("../x.pdf", "x")])
def test_file_names(name, expected):
    assert exporting.stem(name) == expected


# Excel -----------------------------------------------------------------------------------

from openpyxl import load_workbook  # noqa: E402

from app.tables import find_tables  # noqa: E402
from app.textlayout import page_lines  # noqa: E402

WORD = FIXTURES / "word"


def sheets(pdf: bytes):
    return load_workbook(io.BytesIO(exporting.export(pdf, "x.pdf", "xlsx").data)).worksheets


def test_an_arabic_table_reads_right_to_left_in_excel():
    (ws,) = sheets((WORD / "arabic-quotation-word.pdf").read_bytes())
    assert ws.sheet_view.rightToLeft  # its first column on the right, as on the page
    assert [c.value for c in ws[1]] == ["الخدمة", "الكمية", "سعر الوحدة (ريال)", "الإجمالي (ريال)"]
    assert ws["A2"].value == "رسائل تسويقية عبر WhatsApp"
    assert ws["B2"].value == 10000 and ws["D2"].value == 1700.0 and ws["D2"].number_format == "#,##0.00"
    assert ws["A1"].fill.fgColor.rgb == "FFDEEAF6"  # the header's shading
    assert ws["D6"].font.bold  # the total


def test_numbers_become_numbers():
    (ws,) = sheets((WORD / "english-styles-word.pdf").read_bytes())
    assert [c.value for c in ws[1]] == ["Metric", "July", "August", "September"]
    assert ws["B2"].value == 120400 and ws["B3"].value == 0.981 and ws["B3"].number_format == "0.0%"
    assert ws["B4"].value == "55 s"  # not a plain number: stays text
    (ws,) = sheets((WORD / "french-quote-word.pdf").read_bytes())
    assert ws["D2"].value == 2157.6 and ws["D2"].number_format == '#,##0.00 "€"'


def test_a_table_split_by_its_text_is_one_cell():
    # The bilingual contract's Arabic notice line is drawn in three pieces: read right to left.
    page = pymupdf.open(WORD / "bilingual-contract-word.pdf")[0]
    (table,) = find_tables(page, page_lines(page))
    cell = next(c for c in table.cells if c.row == 4 and c.col == 1)
    from app.flow import paragraphs

    assert [p.text for p in paragraphs(cell.lines, cell.bbox[0], cell.bbox[2])] == ["٤. الإشعارات: تُرسل جميع الإشعارات إلى legal@example.com."]


def _grid(merged_header: bool) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    xs, ys = [72, 200, 330, 460], [100, 130, 160, 190]
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y))
    for i, x in enumerate(xs):
        top = ys[1] if merged_header and 0 < i < len(xs) - 1 else ys[0]
        page.draw_line((x, top), (x, ys[-1]))
    page.insert_text((80, 120), "Quarterly totals", fontsize=11)
    for r, y in enumerate((150, 180)):
        for c, x in enumerate(xs[:-1]):
            page.insert_text((x + 8, y), f"{(r + 1) * 100 + c}", fontsize=11)
    return doc.tobytes()


def test_cells_without_a_border_between_them_are_merged():
    page = pymupdf.open("pdf", _grid(merged_header=True))[0]
    (table,) = find_tables(page, page_lines(page))
    header = next(c for c in table.cells if c.row == 0)
    assert (header.rowspan, header.colspan) == (1, 3)
    (ws,) = sheets(_grid(merged_header=True))
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:C1"]
    assert ws["A1"].value == "Quarterly totals" and ws["C3"].value == 202


def test_a_table_without_borders_is_found_by_its_columns():
    doc = pymupdf.open()
    page = doc.new_page()
    rows = [("Item", "Qty", "Price"), ("Paper", "10", "4.50"), ("Ink", "2", "19.90"), ("Pens", "12", "6.00")]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            page.insert_text((72 + c * 150, 100 + r * 20), text, fontsize=11)
    (ws,) = sheets(doc.tobytes())
    assert [[c.value for c in row] for row in ws.iter_rows()] == [["Item", "Qty", "Price"], ["Paper", 10, 4.5], ["Ink", 2, 19.9], ["Pens", 12, 6.0]]


def test_a_pdf_without_tables_gives_its_text():
    ws_list = sheets(make_pdf(2))
    assert [ws.title for ws in ws_list] == ["Page 1", "Page 2"]
    assert ws_list[1]["A1"].value == "Page 2"


@pytest.mark.parametrize(
    "text, value, fmt",
    [
        ("1,700.00", 1700.0, "#,##0.00"),
        ("0.17", 0.17, "0.00"),
        ("98.1%", 0.981, "0.0%"),
        ("89,90 €", 89.9, '0.00 "€"'),
        ("1 450,00 €", 1450.0, '#,##0.00 "€"'),
        ("SAR 4,542.50", 4542.5, '"SAR "#,##0.00'),
        ("12", 12, None),
        ("007", "007", None),
        ("+966 55 123 4567", "+966 55 123 4567", None),
        ("21/10/2026", "21/10/2026", None),
    ],
)
def test_cell_values(text, value, fmt):
    assert exporting._value(text) == (value, fmt)


# PowerPoint with text to edit ------------------------------------------------------------

from pptx.oxml.ns import qn  # noqa: E402

from app.slides import pieces  # noqa: E402


def _deck(name: str):
    return Presentation(io.BytesIO(exporting.export((WORD / f"{name}.pdf").read_bytes(), "x.pdf", "pptx").data))


def _boxes(slide):
    return {s.text_frame.text: s for s in slide.shapes if s.has_text_frame}


def _background(slide) -> Image.Image:
    (picture,) = [s for s in slide.shapes if s.shape_type == 13]
    return Image.open(io.BytesIO(picture.image.blob)).convert("L")


def _ink(image: Image.Image, box, page_width: float) -> int:
    """Dark pixels in a box (page points) of the slide's background picture."""
    k = image.width / page_width
    region = image.crop(tuple(round(v * k) for v in box))
    return sum(1 for v in region.getdata() if v < 128)


def test_every_line_is_a_text_box_over_the_page_without_its_text():
    deck = _deck("english-styles-word")
    page = pymupdf.open(WORD / "english-styles-word.pdf")[0]
    lines = page_lines(page)
    boxes = _boxes(deck.slides[0])
    title = next(l for l in lines if l.text == "Quarterly Service Report")
    shape = boxes["Quarterly Service Report"]
    assert abs(shape.left / 12700 - title.bbox[0]) < 1 and abs(shape.top / 12700 - title.bbox[1]) < 1
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Georgia" and run.font.bold and str(run.font.color.rgb) == title.spans[0][1].color.upper()
    # The background keeps the table's lines and shading, not its text.
    image = _background(deck.slides[0])
    assert _ink(image, title.bbox, page.rect.width) == 0


def test_mixed_arabic_and_english_keep_their_order():
    shape = _boxes(_deck("arabic-quotation-word").slides[0])["رقم العرض: Q-2026-0417"]
    para = shape.text_frame.paragraphs[0]
    assert para._p.pPr.get("rtl") == "1"
    runs = [(r.text, r._r.rPr.get("lang")) for r in para.runs]
    assert runs == [("رقم العرض: ", "ar-SA"), ("Q-2026-0417", "en-US")]
    assert para.runs[0]._r.rPr.find(qn("a:cs")).get("typeface") == "Arial"  # Arabic letters' font


def test_letters_the_pdf_stores_wrongly_stay_in_the_picture():
    page = pymupdf.open(WORD / "arabic-quotation-word.pdf")[0]
    uncertain = [l for l in page_lines(page) if l.uncertain]
    assert uncertain
    slide = _deck("arabic-quotation-word").slides[0]
    assert not any(l.text in _boxes(slide) for l in uncertain)
    assert all(_ink(_background(slide), l.bbox, page.rect.width) > 50 for l in uncertain)


def test_raised_and_lowered_letters():
    page = pymupdf.open(WORD / "english-styles-word.pdf")[0]
    line = next(l for l in page_lines(page) if l.text.startswith("Carbon savings"))
    rises = {text: rise for text, _, _, rise in pieces(line)}
    assert rises["2"] < 0 < rises["rd"]  # CO2, 3rd


# Word ------------------------------------------------------------------------------------

import docx  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _word(pdf: bytes):
    out = exporting.export(pdf, "x.pdf", "docx")
    return docx.Document(io.BytesIO(out.data)), out


def _para(d, start: str):
    return next(p for p in d.paragraphs if p.text.startswith(start))


def test_arabic_paragraphs_are_right_to_left_with_their_fonts():
    d, out = _word((WORD / "arabic-quotation-word.pdf").read_bytes())
    p = _para(d, "رقم العرض")
    assert p.text == "رقم العرض: Q-2026-0417"
    assert p._p.pPr.find(f"{W}bidi") is not None
    arabic, latin = p.runs
    assert arabic._r.rPr.find(f"{W}rtl").get(f"{W}val") is None  # on
    assert latin._r.rPr.find(f"{W}rtl").get(f"{W}val") == "false"
    fonts = arabic._r.rPr.find(f"{W}rFonts")
    assert fonts.get(f"{W}cs") == fonts.get(f"{W}ascii") == "Arial"
    assert arabic._r.rPr.find(f"{W}szCs").get(f"{W}val") == "24"  # 12 pt for Arabic letters too
    # Lines whose letters the PDF stores wrongly are marked, and the person is told.
    assert out.note and "highlighted" in out.note
    assert any(r._r.rPr.find(f"{W}highlight") is not None for p in d.paragraphs for r in p.runs)


def test_a_wrapped_paragraph_is_one_paragraph():
    d, _ = _word((WORD / "english-styles-word.pdf").read_bytes())
    text = next(e["text"] for e in json.loads((WORD / "expected.json").read_text(encoding="utf-8")) if e["text"].startswith("Customers sent more"))
    assert _para(d, "Customers sent more").text == text
    assert _para(d, "Customers sent more")._p.pPr.find(f"{W}jc").get(f"{W}val") == "both"  # justified, as on the page


def test_tables_keep_their_columns_shading_and_direction():
    d, _ = _word((WORD / "arabic-quotation-word.pdf").read_bytes())
    (table,) = d.tables
    assert table._tbl.tblPr.find(f"{W}bidiVisual") is not None  # first column on the right
    assert [c.text for c in table.rows[0].cells] == ["الخدمة", "الكمية", "سعر الوحدة (ريال)", "الإجمالي (ريال)"]
    shading = table.rows[0].cells[0]._tc.tcPr.find(f"{W}shd")
    assert shading.get(f"{W}fill") == "DEEAF6"


def test_links_raised_letters_and_underlines():
    d, _ = _word((WORD / "english-styles-word.pdf").read_bytes())
    p = _para(d, "Full statistics")
    link = p._p.find(f"{W}hyperlink")
    target = p.part.rels[link.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")].target_ref
    assert target == "https://dashboard.example.com/q3"
    assert link.find(f".//{W}u") is not None
    carbon = _para(d, "Carbon savings")
    raised = [r for r in carbon.runs if r._r.rPr.find(f"{W}position") is not None]
    assert {r.text for r in raised} == {"2", "rd"}


def test_headers_footers_and_pictures():
    d, _ = _word((FIXTURES / "images" / "images-word.pdf").read_bytes())
    header = d.sections[0].header
    assert header._element.findall(f".//{W}drawing")  # the logo on every page
    status = _para(d, "Status:")
    kinds = [("picture" if r._r.find(f".//{W}drawing") is not None else r.text) for r in status.runs]
    assert kinds.index("picture") == 2 and kinds[0] == "Status:"  # the icon after "Status: "
    (table,) = d.tables
    assert table.rows[0].cells[1]._tc.findall(f".//{W}drawing")  # the signature in its cell
    d, _ = _word((WORD / "arabic-quotation-word.pdf").read_bytes())
    footer = d.sections[0].footer
    fields = [f.get(f"{W}instr") for f in footer._element.iter(f"{W}fldSimple")]
    assert fields == ["PAGE", "NUMPAGES"]  # "صفحة 1 من 2" counts by itself


def test_each_page_starts_a_page():
    d, _ = _word((WORD / "english-styles-word.pdf").read_bytes())
    breaks = d.element.body.findall(f".//{W}pageBreakBefore")
    assert len(breaks) == 1


def test_odd_pages():
    doc = pymupdf.open()
    doc.new_page()  # empty
    p = doc.new_page()
    p.insert_text((72, 72), "Turned page", fontsize=12)
    p.set_rotation(90)
    p = doc.new_page(width=300, height=200)
    p.insert_image(p.rect, filename=str(FIXTURES / "images" / "photo.jpg"))  # a scan: a picture, no text
    d, _ = _word(doc.tobytes())
    assert [p.text for p in d.paragraphs if p.text] == ["Turned page"]
    assert len(d.inline_shapes) == 1
    assert [(round(s.page_width.pt), round(s.page_height.pt)) for s in d.sections] == [(595, 842), (842, 595), (300, 200)]
