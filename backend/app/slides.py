"""Download as PowerPoint with text you can edit.

Every page becomes a slide in two layers: underneath, a picture of the page with its
text taken out (the same removal that deleting a line uses, so pictures, lines, colors
and shading stay exactly as they were); on top, every line of text in its own text box,
where it was, in its font, size, color and style. Arabic lines read right to left.

Text whose letters the PDF stores wrongly (Line.uncertain), text that isn't drawn
across the page, and invisible text (a scan's hidden text layer) stay in the picture:
in a text box they would show the wrong letters, or twice.
"""

import io

import pikepdf
import pymupdf

from . import exporting
from .changes import rewrite, write_page
from .fonts import family_name
from .pdfcontent import Unsupported, matches_trace, replay
from .textlayout import Line, page_lines

_INVISIBLE = {3, 7}  # render modes that draw nothing


def _editable(line: Line, rp) -> bool:
    if line.uncertain or not line.text.strip():
        return False
    if any(abs(g.direction[1]) > 0.01 or g.direction[0] < 0 for g in line.glyphs):
        return False
    return all(rp.glyphs[t].paint.Tr not in _INVISIBLE for t in line.members)


def _text_box(slide, line: Line, at, scale: float) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    x0, y0, x1, y1 = line.bbox
    ox, oy = at
    # A little room on both sides: the viewer's font may be a hair wider.
    left, top = ox + x0 * scale, oy + y0 * scale
    width, height = (x1 - x0) * scale + 4, (y1 - y0) * scale
    if line.rtl:
        left -= 4
    box = slide.shapes.add_textbox(exporting.emu(left), exporting.emu(top), exporting.emu(width), exporting.emu(height))
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT if line.rtl else PP_ALIGN.LEFT
    para.line_spacing = 1.0
    if line.rtl:
        para._p.get_or_add_pPr().set("rtl", "1")
    for text, style, rtl, rise in pieces(line):
        run = para.add_run()
        run.text = text
        font = run.font
        font.size = Pt(max(1.0, round(style.size * scale * 2) / 2))
        font.bold = style.bold
        font.italic = style.italic
        font.color.rgb = RGBColor.from_string(style.color.upper())
        family = family_name(style.font)
        font.name = family  # Latin letters
        rpr = run._r.get_or_add_rPr()
        # Arabic letters take the "complex script" font: without it PowerPoint draws
        # them in its default font.
        rpr.find(qn("a:latin")).addnext(rpr.makeelement(qn("a:cs"), {"typeface": family}))
        # PowerPoint decides a piece's direction by its language, not by its letters
        # alone: without it, a colon between Arabic and English jumps to the line's end.
        if rtl:
            rpr.set("lang", "ar-SA")
        elif line.rtl:
            rpr.set("lang", "en-US")
        if rise:
            rpr.set("baseline", str(rise))  # raised (3rd) or lowered (CO2), in 1/1000 %


def pieces(line: Line) -> list[tuple[str, object, bool, int]]:
    """The line in typing order, in pieces of one direction, one style and one height:
    (text, style, right-to-left, rise). The rise is how far the piece sits above the
    line's baseline (below, if negative), as PowerPoint and Word measure it: in
    thousandths of a percent of its size."""
    weight: dict = {}
    for g in line.glyphs:
        if not g.is_space and not g.synthetic:
            weight[g.origin[1]] = weight.get(g.origin[1], 0) + 1
    baseline = max(weight, key=weight.get) if weight else 0.0
    out: list[list] = []
    for run in line.runs:
        for i in run.glyphs:
            g = line.glyphs[i]
            text = line.glyph_text(i)
            size = g.style.size or 1.0
            rise = 0 if g.is_space else round((baseline - g.origin[1]) / size * 100) * 1000
            if abs(rise) < 8000:
                rise = 0  # within a point or so: on the line
            if out and out[-1][1] == g.style and out[-1][2] == run.rtl and (out[-1][3] == rise or g.is_space):
                out[-1][0] += text
            else:
                out.append([text, g.style, run.rtl, rise])
    return [tuple(p) for p in out if p[0]]


def deck(src: bytes, name: str) -> exporting.Export:
    doc = pymupdf.open("pdf", src)
    pdf = pikepdf.open(io.BytesIO(src))
    plans = {}
    for n, page in enumerate(doc):
        lines = page_lines(page)
        if page.rotation:
            continue  # a turned page stays one picture
        try:
            rp = replay(pdf, n)
        except Unsupported:
            continue
        if not matches_trace(rp, page) or rp.stray_restore:
            continue
        editable = [l for l in lines if _editable(l, rp)]
        if not editable:
            continue
        write_page(pdf, n, rewrite(rp, {t for l in editable for t in l.members}))
        plans[n] = editable
    out = io.BytesIO()
    pdf.save(out)
    background = pymupdf.open("pdf", out.getvalue())

    slides, size = exporting.new_deck(doc)
    for n, page in enumerate(doc):
        x, y, w, h = exporting._fitted(page, size)
        slide = exporting.blank_slide(slides)
        picture = exporting._picture_for_slide(background[n])
        slide.shapes.add_picture(io.BytesIO(picture), exporting.emu(x), exporting.emu(y), exporting.emu(w), exporting.emu(h))
        scale = w / page.rect.width
        for line in plans.get(n, []):
            _text_box(slide, line, (x, y), scale)
    return exporting.Export(exporting.save_deck(slides), f"{name}.pptx", exporting.PPTX)
