"""Download as…: the edited PDF as another kind of file.

- "jpg", "png": a picture of every page (one file, or a zip of them).
- "pptx-pictures": PowerPoint, every page a slide showing a picture of it: it looks the
  same, but its text can't be edited.

Every format is made from the current PDF, the same one "Download" gives.
"""

import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import pymupdf

DPI = 200  # sharp on screen and in print, without huge files
MAX_PAGE_PIXELS = 25_000_000  # a poster-sized page is drawn at a lower resolution


class UnknownFormat(ValueError):
    pass


@dataclass
class Export:
    data: bytes
    filename: str
    media_type: str


ZIP = "application/zip"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def stem(name: str) -> str:
    """The file's name without .pdf (the name the person's file had)."""
    s = PurePath(name).name
    return s[:-4] if s.lower().endswith(".pdf") else s


def _zoom(page: pymupdf.Page, dpi: int = DPI) -> float:
    zoom = dpi / 72
    pixels = page.rect.width * page.rect.height * zoom * zoom
    if pixels > MAX_PAGE_PIXELS:
        zoom *= (MAX_PAGE_PIXELS / pixels) ** 0.5
    return zoom


def page_picture(page: pymupdf.Page, kind: str, dpi: int = DPI) -> bytes:
    """A page as a picture, as it shows in a PDF viewer (form fields and notes too)."""
    zoom = _zoom(page, dpi)
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    pix.set_dpi(round(zoom * 72), round(zoom * 72))  # so it prints at the page's own size
    return pix.tobytes("jpeg", jpg_quality=90) if kind == "jpg" else pix.tobytes("png")


def _pictures(doc: pymupdf.Document, name: str, kind: str) -> Export:
    media = "image/jpeg" if kind == "jpg" else "image/png"
    if doc.page_count == 1:
        return Export(page_picture(doc[0], kind), f"{name}.{kind}", media)
    digits = len(str(doc.page_count))
    out = io.BytesIO()
    # Pictures are compressed already: storing them is as small and much faster.
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        for n, page in enumerate(doc):
            z.writestr(f"{name} - page {n + 1:0{digits}d}.{kind}", page_picture(page, kind))
    return Export(out.getvalue(), f"{name} (pictures).zip", ZIP)


# PowerPoint ------------------------------------------------------------------------------

EMU_PER_PT = 12700
# PowerPoint's slide size limits: 1 inch to 56 inches.
_MIN_SIDE, _MAX_SIDE = 72, 56 * 72


def slide_size(doc: pymupdf.Document) -> tuple[float, float]:
    """The slides' size, in points: the first page's, within PowerPoint's limits. Pages
    of other sizes are fitted onto it."""
    w, h = doc[0].rect.width, doc[0].rect.height
    scale = min(1.0, _MAX_SIDE / w, _MAX_SIDE / h)
    w, h = w * scale, h * scale
    scale = max(1.0, _MIN_SIDE / w, _MIN_SIDE / h)
    return min(_MAX_SIDE, w * scale), min(_MAX_SIDE, h * scale)


def _fitted(page: pymupdf.Page, slide: tuple[float, float]) -> tuple[float, float, float, float]:
    """Where a page goes on a slide (points): as large as it fits, centered."""
    pw, ph = page.rect.width, page.rect.height
    s = min(slide[0] / pw, slide[1] / ph)
    w, h = pw * s, ph * s
    return (slide[0] - w) / 2, (slide[1] - h) / 2, w, h


def new_deck(doc: pymupdf.Document):
    from pptx import Presentation
    from pptx.util import Emu

    deck = Presentation()
    w, h = slide_size(doc)
    deck.slide_width, deck.slide_height = Emu(round(w * EMU_PER_PT)), Emu(round(h * EMU_PER_PT))
    return deck, (w, h)


def blank_slide(deck):
    return deck.slides.add_slide(deck.slide_layouts[6])  # the blank layout


def emu(points: float):
    from pptx.util import Emu

    return Emu(round(points * EMU_PER_PT))


def save_deck(deck) -> bytes:
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def _picture_for_slide(page: pymupdf.Page) -> bytes:
    """PNG keeps text crisp; a page that's mostly photo is much smaller as JPEG."""
    png = page_picture(page, "png")
    if len(png) < 1_500_000:
        return png
    jpg = page_picture(page, "jpg")
    return jpg if len(jpg) < len(png) else png


def _pptx_pictures(doc: pymupdf.Document, name: str) -> Export:
    deck, size = new_deck(doc)
    for page in doc:
        x, y, w, h = _fitted(page, size)
        slide = blank_slide(deck)
        slide.shapes.add_picture(io.BytesIO(_picture_for_slide(page)), emu(x), emu(y), emu(w), emu(h))
    return Export(save_deck(deck), f"{name}.pptx", PPTX)


# All formats -----------------------------------------------------------------------------

FORMATS = ("jpg", "png", "pptx-pictures")


def export(src: bytes, name: str, fmt: str) -> Export:
    if fmt not in FORMATS:
        raise UnknownFormat(f"Unknown format: {fmt}.")
    doc = pymupdf.open("pdf", src)
    name = stem(name)
    if fmt in ("jpg", "png"):
        return _pictures(doc, name, fmt)
    return _pptx_pictures(doc, name)
